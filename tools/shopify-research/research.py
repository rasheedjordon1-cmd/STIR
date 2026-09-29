#!/usr/bin/env python3
"""Shopify competitor research: scrape stores -> snapshot raw data -> build xlsx report.

    python research.py                       # scrape stores.txt, write snapshot + report
    python research.py --from-snapshot 2026-09-29   # rebuild a report without scraping
"""
from __future__ import annotations

import argparse
import logging
import shutil
import sys
from datetime import date, datetime, timezone
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

import saturation
from common import load_snapshot, product_metrics, read_stores, tags_to_str, write_json
from shopify_client import DEFAULT_UA, PoliteSession, scrape_store

log = logging.getLogger("shopify")

# Flag thresholds
PROVEN_MAX_RANK, PROVEN_MIN_DAYS = 20, 90
TESTING_MAX_DAYS = 21
FAKE_ANCHOR_MIN_PCT = 50

COLUMNS = [
    "store", "title", "url", "vendor", "product_type", "tags", "min_price", "max_price",
    "compare_at", "discount_pct", "variant_count", "pct_variants_available", "image_count",
    "created_at", "days_live", "best_seller_rank",
    # additions beyond the brief
    "PROVEN", "TESTING", "FAKE_ANCHOR", "saturation_stores", "saturation_group",
]
WIDTHS = {"store": 22, "title": 42, "url": 40, "vendor": 18, "product_type": 18, "tags": 30,
          "created_at": 12, "saturation_group": 40}
FILL_PROVEN = PatternFill("solid", fgColor="D8F0D2")
FILL_TESTING = PatternFill("solid", fgColor="FFF2C2")
FILL_ANCHOR = PatternFill("solid", fgColor="F8CBC4")
HEADER_FONT = Font(bold=True, color="FFFFFF")
HEADER_FILL = PatternFill("solid", fgColor="1F1F1F")


# --------------------------------------------------------------------------- scrape

def run_scrape(stores: list[str], snap_dir: Path, args) -> None:
    session = PoliteSession(delay=args.delay, max_retries=args.max_retries, user_agent=args.user_agent)
    started = datetime.now(timezone.utc).isoformat(timespec="seconds")
    summary = []
    for i, domain in enumerate(stores, 1):
        log.info("[%d/%d] %s", i, len(stores), domain)
        res = scrape_store(session, domain, max_pages=args.max_pages, bestseller_pages=args.bestseller_pages)
        d = snap_dir / domain
        shutil.rmtree(d, ignore_errors=True)  # same-day re-run: never mix old and new files
        write_json(d / "store.json", res["store"])
        if res["store"]["status"] == "ok":
            write_json(d / "products.json", res["products"])
            write_json(d / "bestsellers.json", res["bestsellers"])
            if res["meta_text"] is not None:
                (d / "meta.json").write_text(res["meta_text"], encoding="utf-8")
        else:
            log.warning("  skipped: %s — %s", res["store"]["status"], res["store"]["reason"])
        summary.append({k: res["store"][k] for k in ("domain", "status", "product_count", "products_complete")})
    write_json(snap_dir / "manifest.json", {
        "date": snap_dir.name, "started_at": started,
        "finished_at": datetime.now(timezone.utc).isoformat(timespec="seconds"),
        "requests": session.request_count, "stores": summary,
    })
    log.info("Snapshot saved: %s (%d requests)", snap_dir, session.request_count)


# --------------------------------------------------------------------------- rows

def build_rows(snapshot: dict[str, dict], as_of: date) -> list[dict]:
    rows = []
    for domain, s in snapshot.items():
        if s["store"].get("status") != "ok":
            continue
        base = s["store"].get("base_url") or f"https://{domain}"
        rank = {h: i for i, h in enumerate(s["bestsellers"], 1)}
        for p in s["products"]:
            handle = (p.get("handle") or "").lower()
            m = product_metrics(p, as_of)
            r = {
                "store": domain,
                "title": p.get("title") or "",
                "url": f"{base}/products/{p.get('handle', '')}",
                "vendor": p.get("vendor") or "",
                "product_type": p.get("product_type") or "",
                "tags": tags_to_str(p.get("tags")),
                **m,
                "best_seller_rank": rank.get(handle),
            }
            r["PROVEN"] = bool(r["best_seller_rank"] and r["best_seller_rank"] <= PROVEN_MAX_RANK
                               and r["days_live"] is not None and r["days_live"] >= PROVEN_MIN_DAYS)
            r["TESTING"] = r["days_live"] is not None and r["days_live"] <= TESTING_MAX_DAYS
            r["FAKE_ANCHOR"] = r["discount_pct"] is not None and r["discount_pct"] >= FAKE_ANCHOR_MIN_PCT
            rows.append(r)
    return rows


def apply_saturation(rows: list[dict], threshold: int) -> list[dict]:
    if not rows:
        return []
    ids, clusters = saturation.cluster(rows, threshold)
    for r, cid in zip(rows, ids):
        c = clusters[cid]
        r["saturation_stores"] = len(c["stores"])
        r["saturation_group"] = c["label"] if len(c["stores"]) > 1 else ""
    return clusters


# --------------------------------------------------------------------------- xlsx

def _header(ws, cols, widths=None):
    ws.append(cols)
    for i, name in enumerate(cols, 1):
        cell = ws.cell(row=1, column=i)
        cell.font, cell.fill = HEADER_FONT, HEADER_FILL
        cell.alignment = Alignment(vertical="center")
        ws.column_dimensions[get_column_letter(i)].width = (widths or {}).get(name, max(11, len(name) + 2))
    ws.freeze_panes = "A2"


def write_products_sheet(ws, rows):
    _header(ws, COLUMNS, WIDTHS)
    col = {c: i for i, c in enumerate(COLUMNS, 1)}
    rows = sorted(rows, key=lambda r: (r["store"], r["best_seller_rank"] or 10**9, r["days_live"] or 0))
    for r in rows:
        created = r["created_at"]
        vals = dict(r, created_at=created.astimezone(timezone.utc).replace(tzinfo=None) if created else None)
        ws.append([vals.get(c) if not isinstance(vals.get(c), bool) else (vals[c] or None) for c in COLUMNS])
        n = ws.max_row
        ws.cell(n, col["url"]).hyperlink = r["url"]
        for c in ("min_price", "max_price", "compare_at"):
            ws.cell(n, col[c]).number_format = "#,##0.00"
        ws.cell(n, col["created_at"]).number_format = "yyyy-mm-dd"
        fill = FILL_PROVEN if r["PROVEN"] else FILL_TESTING if r["TESTING"] else None
        if fill:
            for i in range(1, len(COLUMNS) + 1):
                ws.cell(n, i).fill = fill
        if r["FAKE_ANCHOR"]:
            ws.cell(n, col["discount_pct"]).fill = FILL_ANCHOR
            ws.cell(n, col["FAKE_ANCHOR"]).fill = FILL_ANCHOR
    ws.auto_filter.ref = ws.dimensions


def write_saturation_sheet(ws, rows, clusters):
    cols = ["product", "stores_selling", "stores", "listings", "min_price", "max_price", "titles_seen"]
    _header(ws, cols, {"product": 44, "stores": 40, "titles_seen": 70})
    shared = [c for c in clusters if len(c["stores"]) > 1]
    shared.sort(key=lambda c: (-len(c["stores"]), -len(c["members"]), c["label"]))
    for c in shared:
        members = [rows[i] for i in c["members"]]
        lows = [m["min_price"] for m in members if m["min_price"] is not None]
        highs = [m["max_price"] for m in members if m["max_price"] is not None]
        titles = list(dict.fromkeys(m["title"] for m in members))
        ws.append([c["label"], len(c["stores"]), ", ".join(sorted(c["stores"])), len(members),
                   min(lows) if lows else None, max(highs) if highs else None, " | ".join(titles[:6])])
    ws.auto_filter.ref = ws.dimensions


def write_stores_sheet(ws, snapshot, rows):
    cols = ["store", "status", "products", "complete", "ranked_bestsellers", "PROVEN", "TESTING",
            "FAKE_ANCHOR", "median_price", "meta_name", "meta_currency", "base_url", "reason"]
    _header(ws, cols, {"store": 24, "status": 16, "base_url": 30, "reason": 60, "meta_name": 24})
    for domain, s in snapshot.items():
        st, meta = s["store"], s["meta"] if isinstance(s["meta"], dict) else {}
        mine = [r for r in rows if r["store"] == domain]
        prices = sorted(r["min_price"] for r in mine if r["min_price"] is not None)
        ws.append([domain, st.get("status"), st.get("product_count"), st.get("products_complete"),
                   st.get("bestseller_count"), sum(r["PROVEN"] for r in mine), sum(r["TESTING"] for r in mine),
                   sum(r["FAKE_ANCHOR"] for r in mine), prices[len(prices) // 2] if prices else None,
                   meta.get("name"), meta.get("currency"), st.get("base_url"), st.get("reason")])


def write_legend_sheet(ws, snap_date):
    lines = [
        ("Snapshot date", snap_date),
        ("PROVEN", f"best_seller_rank <= {PROVEN_MAX_RANK} AND days_live >= {PROVEN_MIN_DAYS}  (green rows)"),
        ("TESTING", f"days_live <= {TESTING_MAX_DAYS}  (yellow rows)"),
        ("FAKE_ANCHOR", f"discount_pct >= {FAKE_ANCHOR_MIN_PCT}  (red cell)"),
        ("best_seller_rank", "Position on /collections/all?sort_by=best-selling, first 3 pages. Blank = not in those pages."),
        ("compare_at", "Highest compare_at_price among variants where compare_at > price."),
        ("discount_pct", "Deepest single-variant markdown: (compare_at - price) / compare_at."),
        ("days_live", "Snapshot date minus created_at (admin creation date, not launch date)."),
        ("saturation_stores", "Number of stores (incl. this one) selling a fuzzy-matched title."),
        ("Prices", "In each store's own currency; see Stores > meta_currency before comparing across stores."),
    ]
    ws.column_dimensions["A"].width, ws.column_dimensions["B"].width = 20, 100
    for k, v in lines:
        ws.append([k, v])
        ws.cell(ws.max_row, 1).font = Font(bold=True)


def build_report(snap_dir: Path, out_path: Path, threshold: int) -> Path:
    snap_date = snap_dir.name
    snapshot = load_snapshot(snap_dir)
    rows = build_rows(snapshot, date.fromisoformat(snap_date))
    clusters = apply_saturation(rows, threshold)
    wb = Workbook()
    write_products_sheet(wb.active, rows)
    wb.active.title = "Products"
    write_saturation_sheet(wb.create_sheet("Saturation"), rows, clusters)
    write_stores_sheet(wb.create_sheet("Stores"), snapshot, rows)
    write_legend_sheet(wb.create_sheet("Legend"), snap_date)
    out_path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(out_path)
    log.info("Report: %s  (%d products, %d stores, %d PROVEN, %d TESTING, %d FAKE_ANCHOR, %d shared products)",
             out_path, len(rows), len(snapshot), sum(r["PROVEN"] for r in rows), sum(r["TESTING"] for r in rows),
             sum(r["FAKE_ANCHOR"] for r in rows), sum(1 for c in clusters if len(c["stores"]) > 1))
    return out_path


# --------------------------------------------------------------------------- cli

def main(argv=None) -> int:
    here = Path(__file__).resolve().parent
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("--stores", type=Path, default=here / "stores.txt")
    ap.add_argument("--snapshots-dir", type=Path, default=here / "snapshots")
    ap.add_argument("--out-dir", type=Path, default=here / "reports")
    ap.add_argument("--date", default=date.today().isoformat(), help="snapshot/report date (default: today)")
    ap.add_argument("--from-snapshot", metavar="YYYY-MM-DD", help="skip scraping; rebuild report from this snapshot")
    ap.add_argument("--delay", type=float, default=1.5, help="seconds between requests (default 1.5)")
    ap.add_argument("--max-retries", type=int, default=3)
    ap.add_argument("--max-pages", type=int, default=100, help="products.json page cap per store (x250 products)")
    ap.add_argument("--bestseller-pages", type=int, default=3)
    ap.add_argument("--sat-threshold", type=int, default=88, help="fuzzy title match score 0-100 (default 88)")
    ap.add_argument("--user-agent", default=DEFAULT_UA)
    ap.add_argument("-v", "--verbose", action="store_true")
    args = ap.parse_args(argv)
    logging.basicConfig(level=logging.DEBUG if args.verbose else logging.INFO,
                        format="%(asctime)s %(levelname)-7s %(message)s", datefmt="%H:%M:%S")
    logging.getLogger("urllib3").setLevel(logging.WARNING)

    snap_date = args.from_snapshot or args.date
    snap_dir = args.snapshots_dir / snap_date
    if not args.from_snapshot:
        stores = read_stores(args.stores)
        if not stores:
            log.error("No stores in %s", args.stores)
            return 1
        run_scrape(stores, snap_dir, args)
    elif not snap_dir.is_dir():
        log.error("No snapshot at %s", snap_dir)
        return 1
    build_report(snap_dir, args.out_dir / f"research_{snap_date}.xlsx", args.sat_threshold)
    return 0


if __name__ == "__main__":
    sys.exit(main())
